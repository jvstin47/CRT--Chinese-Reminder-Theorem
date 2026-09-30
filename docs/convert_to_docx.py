import os
import glob
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def md_to_docx(md_path, docx_path):
    doc = Document()
    
    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_lines = []
    table_lines = []

    def flush_docx_table():
        if not table_lines:
            return
        table_data = []
        for tline in table_lines:
            if '---' in tline or tline == '| | |' or tline == '| |':
                continue
            cols = [c.strip() for c in tline.strip('|').split('|')]
            if len(cols) == 2:
                table_data.append(cols)
        
        if table_data:
            table = doc.add_table(rows=len(table_data), cols=2)
            table.autofit = False
            for r_idx, row_cols in enumerate(table_data):
                row = table.rows[r_idx]
                
                # Format left cell
                p0 = row.cells[0].paragraphs[0]
                p0.paragraph_format.space_before = Pt(2)
                p0.paragraph_format.space_after = Pt(2)
                parse_formatted_text(p0, row_cols[0])
                
                # Format right cell
                p1 = row.cells[1].paragraphs[0]
                p1.paragraph_format.space_before = Pt(2)
                p1.paragraph_format.space_after = Pt(2)
                parse_formatted_text(p1, row_cols[1])
                
                row.cells[0].width = Inches(3.4)
                row.cells[1].width = Inches(3.2)
                
            p_space = doc.add_paragraph()
            p_space.paragraph_format.space_after = Pt(4)
        table_lines.clear()

    for line in lines:
        raw_line = line.rstrip('\n')

        if raw_line.strip().startswith('|'):
            table_lines.append(raw_line.strip())
            continue
        elif table_lines:
            flush_docx_table()
        
        # Explicit Master report page break separator
        if '=========================================================================' in raw_line:
            doc.add_page_break()
            continue

        # Code block toggle
        if raw_line.startswith('```'):
            if in_code_block:
                # End code block
                code_text = '\n'.join(code_lines)
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.left_indent = Inches(0.2)
                p.paragraph_format.keep_with_next = True
                
                # Format code box style
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
                
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
                code_lines = []
            continue

        if in_code_block:
            code_lines.append(raw_line)
            continue

        # Horizontal rule
        if raw_line.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p_border = parse_xml(r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E1"/></w:pBdr>')
            p._p.get_or_add_pPr().append(p_border)
            continue

        # Headings (Times New Roman) with keep_with_next = True
        if raw_line.startswith('# '):
            p = doc.add_heading(level=1)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(raw_line[2:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(19)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Midnight Slate
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(6)
            continue
        elif raw_line.startswith('## '):
            p = doc.add_heading(level=2)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(raw_line[3:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(13.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Academic Navy
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            continue
        elif raw_line.startswith('### '):
            p = doc.add_heading(level=3)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(raw_line[4:])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x33, 0x41, 0x55) # Dark Slate
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            continue

        # Images: ![alt](path) - Preserve natural aspect ratio
        img_match = re.match(r'!\[.*?\]\((.*?)\)', raw_line.strip())
        if img_match:
            img_path = img_match.group(1)
            if os.path.exists(img_path):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                run = p.add_run()
                run.add_picture(img_path, width=Inches(5.5))
            continue

        # Bullet list items
        if raw_line.strip().startswith('- ') or raw_line.strip().startswith('* '):
            item_text = raw_line.strip()[2:]
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            parse_formatted_text(p, item_text)
            continue

        # Numbered list items
        num_match = re.match(r'^\d+\.\s+(.*)', raw_line.strip())
        if num_match:
            item_text = num_match.group(1)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            parse_formatted_text(p, item_text)
            continue

        # Regular Paragraph (Times New Roman)
        if raw_line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            parse_formatted_text(p, raw_line)

    if table_lines:
        flush_docx_table()

    doc.save(docx_path)
    print(f"Successfully created docx with keep_with_next & page breaks: {docx_path}")

def parse_formatted_text(paragraph, text):
    pattern = re.compile(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)')
    tokens = pattern.split(text)
    
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.name = 'Times New Roman'
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Times New Roman'
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)
        else:
            run = paragraph.add_run(token)
            run.font.name = 'Times New Roman'

if __name__ == '__main__':
    docs_dir = '/Users/justin/Public/projects/CRT/docs'
    md_files = glob.glob(os.path.join(docs_dir, '*.md'))
    for md_file in md_files:
        docx_file = md_file[:-3] + '.docx'
        md_to_docx(md_file, docx_file)
